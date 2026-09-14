import re
import io

try:
    import openpyxl
except ImportError:
    openpyxl = None


def _amount(val) -> float:
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(',', '').replace(' ', '').replace('−', '-').replace('–', '-')
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    try:
        return float(s)
    except Exception:
        return 0.0


GROUP_INFO = {
    '100': {'name': 'הכנסות',              'type': 'revenue'},
    '300': {'name': 'הוצאות תפעוליות',     'type': 'opex'},
    '301': {'name': 'הוצאות משאיות',       'type': 'opex'},
    '302': {'name': 'משכורות',             'type': 'salary'},
    '304': {'name': 'עלות דלקים ושמנים',  'type': 'cogs'},
    '305': {'name': 'הוצאות מגרש',        'type': 'opex'},
    '306': {'name': 'הוצאות רכב',         'type': 'opex'},
    '307': {'name': 'ביטוחי רכב ורישיון', 'type': 'opex'},
    '309': {'name': 'הוצאות מימון',       'type': 'finance'},
    '390': {'name': 'הוצאות ביובית',      'type': 'opex'},
    '391': {'name': 'משכורות ביובית',     'type': 'salary'},
    '500': {'name': 'לקוחות',             'type': 'receivable'},
    '501': {'name': 'לקוחות מזדמנים',    'type': 'receivable'},
    '509': {'name': 'לקוחות ביובית',     'type': 'receivable'},
    '600': {'name': 'ספקים',              'type': 'payable'},
    '610': {'name': 'עובדים',             'type': 'other'},
    '700': {'name': 'בנקים',              'type': 'bank'},
    '701': {'name': 'אשראי (חובה)',       'type': 'bank'},
    '702': {'name': 'אשראי (זכות)',       'type': 'bank'},
    '720': {'name': 'בנק מעבר',          'type': 'bank'},
    '750': {'name': 'קופות',             'type': 'cash'},
    '760': {'name': 'קופה קטנה',        'type': 'cash'},
}

# Keyword → (group_code, type) — order matters: more specific first
SUMMARY_KEYWORD_MAP = [
    (['הכנסות'],                                              '100', 'revenue'),
    (['עלות דלקים', 'עלות שמנים', 'עלות מכר', 'עלות המכר'], '304', 'cogs'),
    (['שכר ביובית', 'משכורות ביובית'],                       '391', 'salary'),
    (['משכורות', 'שכר עובדים', 'שכר'],                      '302', 'salary'),
    (['הוצאות מימון', 'עמלות בנק', 'ריבית', 'מימון'],       '309', 'finance'),
    (['הוצאות רכב'],                                         '306', 'opex'),
    (['ביטוח רכב', 'ביטוח'],                                 '307', 'opex'),
    (['הוצאות מגרש', 'מגרש'],                                '305', 'opex'),
    (['הוצאות ביובית'],                                      '390', 'opex'),
    (['הוצאות משאיות', 'משאיות'],                            '301', 'opex'),
    (['הוצאות'],                                             '300', 'opex'),
    (['לקוחות ביובית'],                                      '509', 'receivable'),
    (['לקוחות מזדמנים'],                                     '501', 'receivable'),
    (['לקוחות', 'חייבים'],                                   '500', 'receivable'),
    (['ספקים', 'זכאים'],                                     '600', 'payable'),
    (['בנק', 'עו"ש', 'חשבון עו'],                           '700', 'bank'),
    (['קופה', 'מזומן'],                                      '750', 'cash'),
]


def _load_rows(file_bytes: bytes):
    """Load Excel and return list of rows. Tries read_only first to avoid style errors."""
    bio = io.BytesIO(file_bytes)

    # read_only=True avoids CellStyle/borderID errors
    try:
        wb = openpyxl.load_workbook(bio, data_only=True, read_only=True)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(['' if c is None else c for c in row])
        try:
            wb.close()
        except Exception:
            pass
        return rows
    except Exception:
        pass

    # Fallback: normal mode
    bio.seek(0)
    wb = openpyxl.load_workbook(bio, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(['' if c is None else c for c in row])
    return rows


def _find_header(all_rows):
    """
    Returns (header_idx, is_summary, col_map).
    is_summary=True  → two-column format: name | amount
    is_summary=False → full trial balance: חובה | זכות columns
    """
    col_map = dict(group=None, account=None, name=None,
                   debit=None, credit=None, diff=None, amount=None)

    for i, row in enumerate(all_rows[:40]):
        row_s = [str(c).strip() for c in row]
        full = ' '.join(row_s)

        # Summary format: "מדד" / "תיאור" + "סכום" / "ש"ח" / "יתרה"
        if (('מדד' in full or 'תיאור' in full) and
                ('סכום' in full or 'ש"ח' in full or 'יתרה' in full)):
            for j, s in enumerate(row_s):
                if 'מדד' in s or 'תיאור' in s:
                    col_map['name'] = j
                elif 'סכום' in s or 'ש"ח' in s or 'יתרה' in s:
                    col_map['amount'] = j
            return i + 1, True, col_map

        # Full trial balance
        if (('חובה' in full or 'Debit' in full) and
                ('זכות' in full or 'Credit' in full)):
            for j, s in enumerate(row_s):
                sl = s.lower()
                if s in ('מיון', 'סיווג', 'קבוצה') or sl == 'group':
                    col_map['group'] = j
                elif s in ('חשבון', 'מספר חשבון') or sl in ('account', 'code'):
                    col_map['account'] = j
                elif 'שם' in s or sl in ('name', 'description'):
                    col_map['name'] = j
                elif s in ('חובה', 'Debit'):
                    col_map['debit'] = j
                elif s in ('זכות', 'Credit'):
                    col_map['credit'] = j
                elif 'הפרש' in s or 'יתרה' in s or sl == 'balance':
                    col_map['diff'] = j
            return i + 1, False, col_map

    return 0, False, col_map


def _parse_summary(data_rows, col_name, col_amount):
    groups: dict = {}

    for row in data_rows:
        cells = ['' if c is None else c for c in row]
        if not any(str(c).strip() for c in cells):
            continue

        # Extract name
        name_val = ''
        if col_name is not None and col_name < len(cells):
            name_val = str(cells[col_name]).strip()
        if not name_val:
            for c in cells:
                s = str(c).strip()
                if s and not re.match(r'^[\d,.\-()%\s]+$', s) and len(s) > 2:
                    name_val = s
                    break

        # Extract amount
        amount_val = 0.0
        if col_amount is not None and col_amount < len(cells):
            amount_val = _amount(cells[col_amount])
        if amount_val == 0.0:
            for c in cells:
                v = _amount(c)
                if v != 0.0:
                    amount_val = v
                    break

        if not name_val or amount_val == 0.0:
            continue

        # Match keyword → group
        matched_code = matched_type = matched_gname = None
        for keywords, code, gtype in SUMMARY_KEYWORD_MAP:
            for kw in keywords:
                if kw in name_val:
                    matched_code = code
                    matched_type = gtype
                    matched_gname = GROUP_INFO.get(code, {}).get('name', name_val)
                    break
            if matched_code:
                break

        if not matched_code:
            continue

        if matched_code not in groups:
            groups[matched_code] = {
                'name': matched_gname,
                'type': matched_type,
                'total_debit':  0.0,
                'total_credit': 0.0,
                'total_net':    0.0,
                'accounts':     [],
            }
        val = abs(amount_val)
        groups[matched_code]['total_net'] += val
        groups[matched_code]['accounts'].append({
            'code':   matched_code,
            'name':   name_val,
            'debit':  0.0,
            'credit': val,
            'net':    val,
        })

    return groups


def _parse_full(data_rows, col_map):
    col_group   = col_map['group']
    col_account = col_map['account']
    col_name    = col_map['name']
    col_debit   = col_map['debit']
    col_credit  = col_map['credit']
    col_diff    = col_map['diff']
    groups: dict = {}

    for row in data_rows:
        cells = ['' if c is None else c for c in row]
        row_text = ' '.join(str(c) for c in cells)

        if not any(str(c).strip() for c in cells):
            continue

        # Summary row: "סה"כ לקבוצה: 100"
        m = re.search(r'סה.?כ\s+לקבוצה\s*:?\s*\*{0,2}(\d{3})\*{0,2}', row_text)
        if m:
            gc = m.group(1)
            if gc in groups:
                nums = [_amount(c) for c in cells if _amount(c) != 0.0]
                if len(nums) >= 3:
                    groups[gc]['total_debit']  = nums[-3]
                    groups[gc]['total_credit'] = nums[-2]
                    groups[gc]['total_net']    = nums[-1]
                elif len(nums) == 2:
                    groups[gc]['total_debit']  = nums[0]
                    groups[gc]['total_credit'] = nums[1]
                    groups[gc]['total_net']    = nums[1] - nums[0]
                elif len(nums) == 1:
                    groups[gc]['total_net'] = nums[0]
            continue

        # Data row — find 6-digit account code
        account_code = group_code = None
        if col_account is not None and col_account < len(cells):
            s = str(cells[col_account]).strip()
            if re.match(r'^\d{6}$', s):
                account_code = s
                group_code   = s[:3]
        if col_group is not None and col_group < len(cells):
            s = str(cells[col_group]).strip()
            if re.match(r'^\d{3}$', s):
                group_code = s
        if account_code is None:
            for c in cells:
                s = str(c).strip()
                if re.match(r'^\d{6}$', s):
                    account_code = s
                    group_code   = s[:3]
                    break
        if not account_code:
            continue

        if group_code not in groups:
            info = GROUP_INFO.get(group_code, {})
            groups[group_code] = {
                'name':         info.get('name', f'קבוצה {group_code}'),
                'type':         info.get('type', 'other'),
                'total_debit':  0.0,
                'total_credit': 0.0,
                'total_net':    0.0,
                'accounts':     [],
            }

        name = ''
        if col_name is not None and col_name < len(cells):
            name = str(cells[col_name]).strip()
        if not name:
            for c in cells:
                s = str(c).strip()
                if s and not re.match(r'^[\d,.\-()]+$', s) and len(s) > 2:
                    name = s
                    break

        debit = credit = diff = 0.0
        if col_debit  is not None and col_debit  < len(cells): debit  = _amount(cells[col_debit])
        if col_credit is not None and col_credit < len(cells): credit = _amount(cells[col_credit])
        if col_diff   is not None and col_diff   < len(cells): diff   = _amount(cells[col_diff])
        if diff == 0.0:
            diff = credit - debit
        if debit == 0.0 and credit == 0.0 and diff == 0.0:
            continue

        groups[group_code]['accounts'].append({
            'code':   account_code,
            'name':   name,
            'debit':  debit,
            'credit': credit,
            'net':    diff,
        })

    # Fill totals from accounts if summary row wasn't found
    for gc, gd in groups.items():
        if gd['total_net'] == 0.0 and gd['accounts']:
            gd['total_debit']  = sum(a['debit']  for a in gd['accounts'])
            gd['total_credit'] = sum(a['credit'] for a in gd['accounts'])
            gd['total_net']    = sum(a['net']    for a in gd['accounts'])

    return groups


def parse_trial_balance(file_bytes: bytes) -> dict:
    all_rows = _load_rows(file_bytes)

    debug_rows = [' | '.join(str(c) for c in r[:10]) for r in all_rows[:12]]

    header_idx, is_summary, col_map = _find_header(all_rows)

    # Extract period string
    period = ''
    for row in all_rows[:max(header_idx + 1, 5)]:
        for c in row:
            s = str(c)
            if re.search(r'\d{2}/\d{2}/\d{4}|\d{4}-\d{2}-\d{2}', s):
                period = s.strip()[:120]
                break
        if period:
            break

    data_rows = all_rows[header_idx:] if header_idx > 0 else all_rows[1:]

    if is_summary:
        groups = _parse_summary(data_rows, col_map['name'], col_map['amount'])
    else:
        groups = _parse_full(data_rows, col_map)

    return {
        'period': period,
        'groups': groups,
        '_debug': {
            'header_row':      header_idx,
            'is_summary':      is_summary,
            'cols':            col_map,
            'total_data_rows': len(data_rows),
            'groups_found':    list(groups.keys()),
            'first_rows':      debug_rows,
        },
    }
